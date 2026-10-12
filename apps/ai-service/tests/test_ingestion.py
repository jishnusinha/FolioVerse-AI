import io

import pytest
from corpus import MARKET
from helpers import ingest

from folioverse_ai.events import Emitter
from folioverse_ai.ingestion.chunking import chunk_document
from folioverse_ai.ingestion.parsers import ParseError, parse
from folioverse_ai.ingestion.pipeline import run_ingestion

WS = "0190f0f0-0000-7000-8000-00000000000a"


def make_pdf() -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    c.setFont("Helvetica-Bold", 20)
    c.drawString(72, 780, "Quarterly Review")
    c.setFont("Helvetica", 11)
    c.drawString(72, 740, "Revenue grew 12 percent in the northern region during the quarter.")
    c.showPage()
    c.setFont("Helvetica", 11)
    c.drawString(72, 780, "Operating costs fell because logistics contracts were renegotiated.")
    c.save()
    return buf.getvalue()


def make_docx() -> bytes:
    from docx import Document

    d = Document()
    d.add_heading("Field Notes", level=1)
    d.add_paragraph("The survey team recorded forty species along the river bank.")
    d.add_heading("Methods", level=2)
    d.add_paragraph("Transects were walked at dawn on three consecutive days.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def make_epub() -> bytes:
    from ebooklib import epub

    book = epub.EpubBook()
    book.set_identifier("demo-book")
    book.set_title("Demo Book")
    book.set_language("en")
    ch1 = epub.EpubHtml(
        title="One", file_name="ch1.xhtml", content="<h1>One</h1><p>Mara crossed the bridge at dawn.</p>"
    )
    ch2 = epub.EpubHtml(
        title="Two", file_name="ch2.xhtml", content="<h1>Two</h1><p>Tomas waited for the ferry.</p>"
    )
    for ch in (ch1, ch2):
        book.add_item(ch)
    book.toc = (ch1, ch2)
    book.spine = ["nav", ch1, ch2]
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())
    buf = io.BytesIO()
    epub.write_epub(buf, book)
    return buf.getvalue()


def test_markdown_headings_and_pages():
    doc = parse(MARKET.encode(), "text/markdown")
    assert doc.pages and doc.elements
    risk = next(e for e in doc.elements if "Supply chain" in e.text)
    assert risk.heading_path[-1] == "Risks"
    page = doc.pages[risk.page - 1]
    assert page.text[risk.start : risk.end] == risk.text


def test_pdf_pages_with_bbox():
    doc = parse(make_pdf(), "application/pdf")
    assert len(doc.pages) == 2
    el = next(e for e in doc.elements if "Revenue grew" in e.text)
    assert el.page == 1 and el.bbox and len(el.bbox) == 4
    assert any(e.kind == "heading" for e in doc.elements)


def test_docx_headings():
    doc = parse(make_docx(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    el = next(e for e in doc.elements if "Transects" in e.text)
    assert el.heading_path == ["Field Notes", "Methods"]


def test_epub_spine_order_and_href():
    doc = parse(make_epub(), "application/epub+zip")
    texts = [e.text for e in doc.elements]
    assert texts.index(next(t for t in texts if "Mara" in t)) < texts.index(
        next(t for t in texts if "Tomas" in t)
    )
    assert any(e.href and e.href.startswith("ch1.xhtml") for e in doc.elements)


def test_unsupported_and_corrupt():
    with pytest.raises(ParseError):
        parse(b"%PDF-1.7 garbage", "application/pdf")
    with pytest.raises(ParseError):
        parse(b"abc", "application/x-msdownload")


def test_chunk_ids_are_deterministic():
    doc = parse(MARKET.encode(), "text/markdown")
    a = chunk_document(doc, "0190f0f0-0000-7000-8000-0000000000aa")
    b = chunk_document(parse(MARKET.encode(), "text/markdown"), "0190f0f0-0000-7000-8000-0000000000aa")
    assert [c["id"] for c in a] == [c["id"] for c in b]
    children = [c for c in a if c["level"] == "CHILD"]
    assert children and all(c["parent_id"] for c in children)


async def test_pipeline_end_to_end(repo):
    svid = await ingest(MARKET, WS)
    assert any(c["source_version_id"] == svid for c in repo.chunks.values())
    assert repo.manifests[svid]["child_count"] > 0
    usage = [u for u in repo.usage_events if u.get("cost_class") == "LOCAL_COMPUTE"]
    assert usage and usage[0]["model_config_id"] == "local-compute"


async def test_reingest_is_idempotent(repo):
    svid = await ingest(MARKET, WS)
    before = sorted(k for k, v in repo.chunks.items() if v["source_version_id"] == svid)
    em = Emitter(None, WS)
    await run_ingestion(
        {
            "jobId": "j",
            "workspaceId": WS,
            "sourceId": "s",
            "sourceVersionId": svid,
            "mimeType": "text/markdown",
            "objectUrl": "x",
        },
        em,
        data=MARKET.encode(),
    )
    after = sorted(k for k, v in repo.chunks.items() if v["source_version_id"] == svid)
    assert before == after


async def test_empty_text_fails_without_raising():
    em = Emitter(None, WS)
    res = await run_ingestion(
        {
            "jobId": "j",
            "workspaceId": WS,
            "sourceId": "s",
            "sourceVersionId": "v",
            "mimeType": "text/plain",
            "objectUrl": "x",
        },
        em,
        data=b"   \n\n  ",
    )
    assert res["state"] == "FAILED"
    assert em.history[-1]["type"] == "ingestion.finished"
