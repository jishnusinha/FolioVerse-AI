"""Format parsers producing a normalized element model.

Every element carries a locator: page number, page-relative character offsets into the page
text (stored in ai.source_pages), optional bbox in PDF points, heading path, and for EPUB the
chapter href. Pages that fail to parse are reported so the version can be marked PARTIAL.
"""

import io
import re
from dataclasses import dataclass, field

SYNTHETIC_PAGE_CHARS = 3000


@dataclass
class Element:
    page: int
    text: str
    start: int
    end: int
    kind: str = "paragraph"
    bbox: list[float] | None = None
    heading_path: list[str] = field(default_factory=list)
    href: str | None = None


@dataclass
class Page:
    page_no: int
    text: str
    width: float | None = None
    height: float | None = None


@dataclass
class ParsedDocument:
    pages: list[Page]
    elements: list[Element]
    failed_pages: list[int] = field(default_factory=list)
    title: str | None = None


class ParseError(Exception):
    pass


def _pages_from_blocks(blocks: list[tuple[str, str, list[str], str | None]]) -> ParsedDocument:
    """Groups (kind, text, heading_path, href) blocks into synthetic pages of ~3000 chars."""
    pages: list[Page] = []
    elements: list[Element] = []
    buf = ""
    page_no = 1
    for kind, text, path, href in blocks:
        text = text.strip()
        if not text:
            continue
        if buf and len(buf) + len(text) > SYNTHETIC_PAGE_CHARS:
            pages.append(Page(page_no, buf))
            page_no += 1
            buf = ""
        start = len(buf) + (2 if buf else 0)
        buf = f"{buf}\n\n{text}" if buf else text
        elements.append(Element(page_no, text, start, start + len(text), kind, None, list(path), href))
    if buf:
        pages.append(Page(page_no, buf))
    return ParsedDocument(pages, elements)


def parse_text(data: bytes, markdown: bool) -> ParsedDocument:
    text = data.decode("utf-8", errors="strict").replace("\r\n", "\n")
    blocks = []
    path: list[str] = []
    for para in re.split(r"\n\s*\n", text):
        para = para.strip()
        if not para:
            continue
        m = re.match(r"^(#{1,6})\s+(.*)$", para) if markdown else None
        if m:
            level = len(m.group(1))
            path = path[: level - 1] + [m.group(2).strip()]
            blocks.append(("heading", m.group(2).strip(), path[:-1], None))
        else:
            blocks.append(("paragraph", re.sub(r"\s*\n\s*", " ", para), list(path), None))
    return _pages_from_blocks(blocks)


def parse_docx(data: bytes) -> ParsedDocument:
    import docx

    d = docx.Document(io.BytesIO(data))
    blocks = []
    path: list[str] = []
    for p in d.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        style = (p.style.name or "").lower() if p.style is not None else ""
        m = re.match(r"heading\s*(\d)", style)
        if m or style == "title":
            level = int(m.group(1)) if m else 1
            path = path[: level - 1] + [text]
            blocks.append(("heading", text, path[:-1], None))
        else:
            blocks.append(("paragraph", text, list(path), None))
    for t in d.tables:
        rows = [" | ".join(c.text.strip() for c in r.cells) for r in t.rows]
        if rows:
            blocks.append(("table", "\n".join(rows), list(path), None))
    return _pages_from_blocks(blocks)


def parse_pdf(data: bytes) -> ParsedDocument:
    import pdfplumber

    pages: list[Page] = []
    elements: list[Element] = []
    failed: list[int] = []
    try:
        pdf = pdfplumber.open(io.BytesIO(data))
    except Exception as e:
        raise ParseError(f"Unreadable PDF: {e.__class__.__name__}") from e
    with pdf:
        title = (pdf.metadata or {}).get("Title")
        for i, page in enumerate(pdf.pages, start=1):
            try:
                words = page.extract_words(keep_blank_chars=False, use_text_flow=True, extra_attrs=["size"])
            except Exception:
                failed.append(i)
                pages.append(Page(i, "", float(page.width), float(page.height)))
                continue
            lines: list[list[dict]] = []
            for w in words:
                if lines and abs(lines[-1][-1]["top"] - w["top"]) < 3:
                    lines[-1].append(w)
                else:
                    lines.append([w])
            # Group lines into paragraph blocks by vertical gap.
            blocks: list[list[list[dict]]] = []
            for ln in lines:
                if blocks:
                    prev = blocks[-1][-1]
                    gap = ln[0]["top"] - prev[0]["bottom"]
                    height = max(1.0, prev[0]["bottom"] - prev[0]["top"])
                    if gap < height * 0.9:
                        blocks[-1].append(ln)
                        continue
                blocks.append([ln])
            page_text = ""
            for b in blocks:
                text = " ".join(" ".join(w["text"] for w in ln) for ln in b).strip()
                if not text:
                    continue
                x0 = min(w["x0"] for ln in b for w in ln)
                x1 = max(w["x1"] for ln in b for w in ln)
                top = min(w["top"] for ln in b for w in ln)
                bottom = max(w["bottom"] for ln in b for w in ln)
                size = max(w.get("size", 0) for ln in b for w in ln)
                start = len(page_text) + (2 if page_text else 0)
                page_text = f"{page_text}\n\n{text}" if page_text else text
                kind = "heading" if size >= 14 and len(text) < 120 else "paragraph"
                elements.append(
                    Element(
                        i,
                        text,
                        start,
                        start + len(text),
                        kind,
                        [round(x0, 2), round(top, 2), round(x1, 2), round(bottom, 2)],
                    )
                )
            pages.append(Page(i, page_text, float(page.width), float(page.height)))
    # Heading paths from large-font blocks.
    path: list[str] = []
    for el in elements:
        if el.kind == "heading":
            path = [el.text]
        else:
            el.heading_path = list(path)
    return ParsedDocument(pages, elements, failed, title)


def parse_epub(data: bytes) -> ParsedDocument:
    import ebooklib
    from ebooklib import epub
    from lxml import html as lhtml

    try:
        book = epub.read_epub(io.BytesIO(data), options={"ignore_ncx": True})
    except Exception as e:
        raise ParseError(f"Unreadable EPUB: {e.__class__.__name__}") from e
    pages: list[Page] = []
    elements: list[Element] = []
    failed: list[int] = []
    titles = book.get_metadata("DC", "title")
    page_no = 0
    for idref, _linear in book.spine:
        item = book.get_item_with_id(idref)
        if item is None or item.get_type() != ebooklib.ITEM_DOCUMENT:
            continue
        page_no += 1
        try:
            doc = lhtml.fromstring(item.get_content())
        except Exception:
            failed.append(page_no)
            pages.append(Page(page_no, ""))
            continue
        chapter_title = None
        for h in doc.iter("h1", "h2", "h3"):
            chapter_title = " ".join(h.text_content().split())
            break
        text = ""
        for n, el in enumerate(doc.iter("p", "h1", "h2", "h3", "li", "blockquote")):
            t = " ".join(el.text_content().split())
            if not t:
                continue
            start = len(text) + (2 if text else 0)
            text = f"{text}\n\n{t}" if text else t
            kind = "heading" if el.tag in ("h1", "h2", "h3") else "paragraph"
            elements.append(
                Element(
                    page_no,
                    t,
                    start,
                    start + len(t),
                    kind,
                    None,
                    [chapter_title] if chapter_title else [],
                    f"{item.get_name()}#p{n}",
                )
            )
        pages.append(Page(page_no, text))
    return ParsedDocument(pages, elements, failed, titles[0][0] if titles else None)


def parse(data: bytes, mime: str) -> ParsedDocument:
    if mime == "application/pdf":
        return parse_pdf(data)
    if mime.endswith("wordprocessingml.document"):
        return parse_docx(data)
    if mime == "application/epub+zip":
        return parse_epub(data)
    if mime == "text/markdown":
        return parse_text(data, markdown=True)
    if mime.startswith("text/"):
        return parse_text(data, markdown=False)
    raise ParseError(f"Unsupported type {mime}")
