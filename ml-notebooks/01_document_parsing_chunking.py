# ---
# jupyter:
#   jupytext:
#     formats: ipynb,py:percent
#     text_representation:
#       extension: .py
#       format_name: percent
#   kernelspec:
#     display_name: Python 3
#     language: python
#     name: python3
# ---

# %% [markdown]
# # 1 · Document parsing & chunking
#
# How well does the FolioVerse ingestion front end turn files into citable units?
#
# * **Parsers** (`folioverse_ai.ingestion.parsers`): PDF (pdfplumber, with boxes), EPUB (chapters, CFIs),
#   DOCX (headings, tables), Markdown and plain text, all into one element model with page-relative offsets.
# * **Chunking** (`folioverse_ai.ingestion.chunking`): ~1200-token parents for the model's context and
#   ~300-token children (15 % overlap) for retrieval, with locators for citation anchors.
#
# Measured here: element and locator coverage per format, offset accuracy (does `page.text[start:end]`
# reproduce the element?), chunk size distribution and overlap, whether answer evidence survives chunking
# intact, and a child-size sweep. Data: `datasets/synthetic` only (fictional DEMO documents).

# %%
import os
import sys

sys.path.insert(0, os.environ.get("FV_NOTEBOOK_DIR", os.getcwd()))
import fvnb

fvnb.setup()
print("smoke mode" if fvnb.SMOKE else "full mode", "-", fvnb.ROOT.name)

# %% [markdown]
# ## Corpus
# Two committed synthetic files plus three generated from catalog descriptions, so every parser is covered.

# %%
import io

import docx

catalog = fvnb.catalog()[: fvnb.sample(240, 40)]
by_genre: dict[str, list[dict]] = {}
for it in catalog:
    by_genre.setdefault(it["genres"][0], []).append(it)

md = "# Synthetic catalog notes (DEMO)\n\n" + "\n\n".join(
    f"## {genre}\n\n" + "\n\n".join(f"### {it['title']}\n\n{it['description']}" for it in items)
    for genre, items in sorted(by_genre.items())
)
plain = "\n\n".join(it["description"] for it in catalog)

d = docx.Document()
d.add_heading("Synthetic catalog notes (DEMO)", level=1)
for genre, items in sorted(by_genre.items()):
    d.add_heading(genre, level=2)
    for it in items:
        d.add_heading(it["title"], level=3)
        d.add_paragraph(it["description"])
table = d.add_table(rows=1, cols=3)
table.rows[0].cells[0].text, table.rows[0].cells[1].text, table.rows[0].cells[2].text = (
    "Title",
    "Year",
    "Genre",
)
for it in catalog[:10]:
    row = table.add_row().cells
    row[0].text, row[1].text, row[2].text = it["title"], str(it["year"]), it["genres"][0]
buf = io.BytesIO()
d.save(buf)

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
corpus = {
    "atlas-brief.pdf": ((fvnb.DATA / "synthetic/docs/atlas-brief.pdf").read_bytes(), "application/pdf"),
    "small-lanterns.epub": (
        (fvnb.DATA / "synthetic/books/small-lanterns.epub").read_bytes(),
        "application/epub+zip",
    ),
    "catalog-notes.docx": (buf.getvalue(), DOCX),
    "catalog-notes.md": (md.encode(), "text/markdown"),
    "catalog-notes.txt": (plain.encode(), "text/plain"),
}
{name: f"{len(data):,} bytes" for name, (data, _) in corpus.items()}

# %% [markdown]
# ## Parsing: coverage and locator accuracy

# %%
from collections import Counter

from folioverse_ai.ingestion.parsers import parse

parsed, parse_rows = {}, []
for name, (data, mime) in corpus.items():
    with fvnb.Timer() as t:
        doc = parse(data, mime)
    parsed[name] = doc
    pages = {p.page_no: p.text for p in doc.pages}
    exact = sum(pages.get(e.page, "")[e.start : e.end] == e.text for e in doc.elements)
    kinds = Counter(e.kind for e in doc.elements)
    n = len(doc.elements) or 1
    parse_rows.append(
        {
            "file": name,
            "pages": len(doc.pages),
            "elements": len(doc.elements),
            "kinds": ", ".join(f"{k} {v}" for k, v in kinds.most_common()),
            "offset_exact": exact / n,
            "with_heading": sum(bool(e.heading_path) for e in doc.elements) / n,
            "with_bbox": sum(e.bbox is not None for e in doc.elements) / n,
            "with_cfi": sum(e.cfi is not None for e in doc.elements) / n,
            "failed_pages": len(doc.failed_pages),
            "parse_ms": t.seconds * 1000,
        }
    )
fvnb.table(parse_rows)

# %% [markdown]
# `offset_exact` is the share of elements whose page-relative offsets reproduce their text exactly — what
# citation highlighting relies on. Boxes exist only for PDF and CFIs only for EPUB, as designed.

# %% [markdown]
# ## Chunking: sizes, overlap and locators

# %%
from folioverse_ai.ingestion import chunking
from folioverse_ai.text import count_tokens, tokens

chunks = {name: chunking.chunk_document(doc, f"sv-{i}") for i, (name, doc) in enumerate(parsed.items())}


def overlap_ratio(children: list[dict]) -> float:
    """Share of each child's leading tokens repeated from the previous child of the same parent."""
    ratios = []
    for a, b in zip(children, children[1:], strict=False):
        if a["parent_id"] != b["parent_id"]:
            continue
        ta, tb = tokens(a["text"]), tokens(b["text"])
        k = next((n for n in range(min(len(ta), len(tb)), 0, -1) if ta[-n:] == tb[:n]), 0)
        ratios.append(k / max(1, len(tb)))
    return fvnb.mean(ratios)


chunk_rows = []
for name, cs in chunks.items():
    kids = [c for c in cs if c["level"] == "CHILD"]
    sizes = [c["token_count"] for c in kids]
    doc = parsed[name]
    chunk_rows.append(
        {
            "file": name,
            "parents": sum(c["level"] == "PARENT" for c in cs),
            "children": len(kids),
            "child_p50": fvnb.percentile(sizes, 0.5),
            "child_p95": fvnb.percentile(sizes, 0.95),
            "child_max": max(sizes, default=0),
            "overlap": overlap_ratio(kids),
            "locator_page": fvnb.mean(c["locator"].get("page") is not None for c in kids),
            "cross_page": fvnb.mean(c["page_start"] != c["page_end"] for c in kids),
            "page_breaks_enforced": doc.break_pages,
        }
    )
fvnb.table(chunk_rows)

# %% [markdown]
# EPUB chapters are pages and `break_pages` is on, so `cross_page` must be 0 there: a chunk never spans two
# chapters (the reader's spoiler guard depends on it). The child cap is 300 tokens; elements longer than that
# are split before grouping, overlap can push a child slightly above the cap.

# %%
import matplotlib.pyplot as plt

fig, ax = plt.subplots(figsize=(7, 3))
for name, cs in chunks.items():
    sizes = [c["token_count"] for c in cs if c["level"] == "CHILD"]
    if len(sizes) > 1:
        ax.hist(sizes, bins=20, alpha=0.5, label=name)
ax.axvline(chunking.CHILD_TOKENS, color="k", linestyle="--", linewidth=1)
ax.set_xlabel("child chunk tokens")
ax.set_ylabel("chunks")
ax.legend(fontsize=7)
fig.tight_layout()

# %% [markdown]
# ## Evidence integrity
# For each answerable question in `qa-small`, is the evidence string wholly inside at least one child chunk
# of the brief? A split evidence span can only be cited partially.

# %%
qa = [q for q in fvnb.qa_small() if q["answerable"]]
pdf_children = [c for c in chunks["atlas-brief.pdf"] if c["level"] == "CHILD"]
integrity = fvnb.mean(any(q["evidence"].lower() in c["text"].lower() for c in pdf_children) for q in qa)
print(f"evidence wholly inside one child chunk: {integrity:.0%} of {len(qa)} questions")

# %% [markdown]
# ## Child size sweep
# Smaller children retrieve more precisely but fragment context; larger children dilute embeddings. The
# sweep re-chunks the Markdown and PDF documents with different child budgets (parents unchanged).

# %%
sweep = []
default = chunking.CHILD_TOKENS
try:
    for size in [80, 150, 300, 450] if not fvnb.SMOKE else [150, 300]:
        chunking.CHILD_TOKENS = size
        md_kids = [
            c for c in chunking.chunk_document(parsed["catalog-notes.md"], "sv-md") if c["level"] == "CHILD"
        ]
        pdf_kids = [
            c for c in chunking.chunk_document(parsed["atlas-brief.pdf"], "sv-pdf") if c["level"] == "CHILD"
        ]
        sweep.append(
            {
                "child_tokens": size,
                "md_children": len(md_kids),
                "md_mean_tokens": fvnb.mean(c["token_count"] for c in md_kids),
                "pdf_children": len(pdf_kids),
                "evidence_intact": fvnb.mean(
                    any(q["evidence"].lower() in c["text"].lower() for c in pdf_kids) for q in qa
                ),
            }
        )
finally:
    chunking.CHILD_TOKENS = default
fvnb.table(sweep)

# %% [markdown]
# ## Findings

# %%
worst = min(parse_rows, key=lambda r: r["offset_exact"])
epub = next(r for r in chunk_rows if r["file"].endswith(".epub"))
fvnb.finding(f"Lowest offset accuracy: {worst['file']} at {worst['offset_exact']:.1%} of elements.")
fvnb.finding(f"Evidence intact at the served child size ({chunking.CHILD_TOKENS} tokens): {integrity:.0%}.")
fvnb.finding(f"EPUB chunks crossing a chapter boundary: {epub['cross_page']:.0%} (must be 0).")
fvnb.finding(
    f"Largest child: {max(r['child_max'] for r in chunk_rows)} tokens (cap {chunking.CHILD_TOKENS} plus overlap)."
)

# %%
fvnb.save(
    "01-parsing-chunking",
    {
        "parsing": parse_rows,
        "chunking": chunk_rows,
        "evidenceIntegrity": integrity,
        "childSizeSweep": sweep,
        "tokenizer": f"approximate ({count_tokens.__module__}.count_tokens)",
    },
)
